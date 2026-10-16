"""Extract Slovak BT dictionary and Slovak rule texts from the FS SR workbook
'Transpozície štandardu Peppol BIS v podmienkach Slovenskej legislatívy' (v1.12).
Usage: python extract_fs.py <workbook.xlsx> <out_dir>
"""
import json, os, re, sys
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
out = sys.argv[2]
os.makedirs(out, exist_ok=True)


def clean(v):
    if v is None:
        return ''
    s = str(v).replace('\r', '').strip()
    s = re.sub(r'[ \t]+', ' ', s)
    s = re.sub(r'\n\s*\n+', '\n', s)
    return s


# --- BT dictionary (sheet "Obchod. term.(legis.) validácia") ---
ws = wb.worksheets[1]
cols = {c: i for i, c in enumerate('ABCDEFGHIJKLMNOPQRSTUVWXYZ')}
bt = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    row = list(row) + [None] * 26
    code = clean(row[cols['M']])
    if not re.match(r'^(BT|BG)-\d+(-\d)?$', code):
        continue
    code = re.sub(r'^(B[TG])-0+(\d)', r'\1-\2', code)   # BG-01 → BG-1
    name = clean(row[cols['S']])
    entry = bt.setdefault(code, {'n': '', 'd': '', 'u': [], 'r': '', 'x': [], 'law': '', 'url': ''})
    if not entry['n'] and name:
        entry['n'] = name
    desc = clean(row[cols['T']])
    if desc and not entry['d']:
        entry['d'] = desc
    use = clean(row[cols['U']])
    if use and use not in entry['u']:
        entry['u'].append(use)
    req = clean(row[cols['R']])
    if req and not entry['r']:
        entry['r'] = req
    for c in ('X', 'Y'):
        t = clean(row[cols[c]])
        if t and t not in entry['x']:
            entry['x'].append(t)
    law = ' '.join(x for x in (clean(row[cols['C']]), clean(row[cols['D']])) if x)
    if law and not entry['law']:
        entry['law'] = re.sub(r'\s+', ' ', law)
    url = clean(row[cols['N']]).split(',')[0].strip()
    if url.startswith('http') and not entry['url']:
        entry['url'] = url

# názvy, ktoré vo workbooku chýbajú (stĺpec S je prázdny)
FALLBACK = {'BT-5': 'Kód meny faktúry', 'BT-157-1': 'Identifikátor schémy'}
for code, name in FALLBACK.items():
    bt.setdefault(code, {'n': '', 'd': '', 'u': [], 'r': '', 'x': [], 'law': '', 'url': ''})
    if not bt[code]['n']:
        bt[code]['n'] = name

for e in bt.values():
    e['u'] = '\n'.join(e['u'])
    e['x'] = '\n'.join(e['x'])
    for k in list(e):
        if not e[k]:
            del e[k]

# --- Slovak rule texts (sheet "UBL Pravidlá") ---
ws = wb.worksheets[2]
rules = {}
for row in ws.iter_rows(min_row=4, values_only=True):
    rid, txt = clean(row[0]), clean(row[1]) if len(row) > 1 else ''
    if not rid or not txt or not re.match(r'^[A-Z]', rid):
        continue
    txt = re.sub(r'^\[' + re.escape(rid) + r'\]\s*-?\s*', '', txt)
    rules[rid] = txt

json.dump(bt, open(os.path.join(out, 'sk_bt.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
json.dump(rules, open(os.path.join(out, 'sk_rules.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)

# --- single-byte code page tables for mojibake detection ---
tables = {}
for cp in ('cp1250', 'cp1252'):
    t = []
    for b in range(0x80, 0x100):
        try:
            t.append(ord(bytes([b]).decode(cp)))
        except UnicodeDecodeError:
            t.append(b)  # undefined byte: lenient converters keep the C1 code point
    tables[cp] = t
json.dump(tables, open(os.path.join(out, 'codepages.json'), 'w', encoding='utf-8'))

print(len(bt), 'BT/BG entries;', len(rules), 'rules')
print('sample:', json.dumps(bt.get('BT-86'), ensure_ascii=False)[:400])
print('sample rule:', rules.get('BR-CO-15'))
